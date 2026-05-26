from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

output_path = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"

headers = [
    "Modul",
    "Informationspunkt",
    "Pflichtfelder (mindestens)",
    "URL 1",
    "URL 2",
    "URL 3",
    "URL 4",
    "Abgedeckter Zeitraum",
    "Bemerkungen",
]

rows = [
    [
        "Airline-Stammdaten",
        "Airline-Liste und Codes",
        "Airline-Name, IATA/ICAO, Land/Hub-Flughafen",
        "", "", "", "", "", "",
    ],
    [
        "Strecken- und Frequenzangebot",
        "Strecken und Frequenzen aller Airlines",
        "Abflugflughafen, Zielflughafen, Fluege pro Woche, Abflugzeitfenster, Flugzeugtyp (falls verfuegbar)",
        "", "", "", "", "", "",
    ],
    [
        "Sitzplatz-/Kapazitaetsangebot",
        "Kapazitaetsdaten aller Airlines",
        "Strecke, Flugzeugtyp, Sitzplatzanzahl (oder ableitbare Felder)",
        "", "", "", "", "", "",
    ],
    [
        "Preisdaten",
        "Preisniveau je Strecke/Zeitslot",
        "Crawling-Datum, Abflugdatum, Airline, Strecke, Mindestpreis/Medianpreis, Tarif-/Klassenhinweis",
        "", "", "", "", "", "",
    ],
    [
        "Nachfrage-Proxies",
        "Such- und Aufmerksamkeitsindikatoren",
        "Keywords, Region, Zeitgranularitaet, Trendindex",
        "", "", "", "", "", "",
    ],
    [
        "Operative Qualitaet",
        "Puenktlichkeit und Annullierungen",
        "Airline/Strecke, OTP, Cancellation Rate, Delay-Verteilung, Definitionsbasis",
        "", "", "", "", "", "",
    ],
    [
        "Kalender- und Eventdaten",
        "Exogene Nachfragefaktoren",
        "Feiertage, Messen, Sportevents, Datum, Stadt",
        "", "", "", "", "", "",
    ],
    [
        "Neue-Strecken-Potenzial",
        "O&D-Nachfragepotenzial",
        "Staedtepaar, Bevoelkerungs-/Tourismus-/Business-Indikatoren, Saisonalitaet",
        "", "", "", "", "", "",
    ],
    [
        "Wettbewerbsstruktur",
        "Wettbewerbslage pro Strecke",
        "Anzahl Carrier pro Strecke, Gesamtfrequenz/Gesamtsitze, Preisband",
        "", "", "", "", "", "",
    ],
    [
        "Flughafen-Restriktionen",
        "Operative und Slot-bezogene Grenzen",
        "Slot-Lage, Nachtflugbeschraenkungen, Terminal/Ground-Handling-Limits, regulatorische Auflagen",
        "", "", "", "", "", "",
    ],
]

wb = Workbook()
ws = wb.active
ws.title = "Datenquellen"

# Write headers
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="00549F", end_color="00549F", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Write data rows
for row_idx, row_data in enumerate(rows, start=2):
    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# Styling borders
thin = Side(style="thin", color="D9D9D9")
for r in range(1, len(rows) + 2):
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Column widths
widths = {
    1: 24,
    2: 34,
    3: 58,
    4: 30,
    5: 30,
    6: 30,
    7: 30,
    8: 22,
    9: 30,
}
for col_idx, width in widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# Row heights
ws.row_dimensions[1].height = 28
for r in range(2, len(rows) + 2):
    ws.row_dimensions[r].height = 52

# Freeze header
ws.freeze_panes = "A2"

# Auto filter
ws.auto_filter.ref = f"A1:I{len(rows) + 1}"

# Add instruction sheet
ins = wb.create_sheet(title="Hinweise")
ins["A1"] = "Hinweise zur Befuellung"
ins["A1"].font = Font(bold=True, size=13, color="00549F")
ins["A3"] = "1) Pro Informationspunkt bitte mehrere URLs eintragen (empfohlen: mindestens 2)."
ins["A4"] = "2) Datensammlung erfolgt fuer alle Airlines ohne Vorab-Trennung in Zielairline vs. Wettbewerber."
ins["A5"] = "3) In Bemerkungen bitte angeben, ob Felder direkt crawlbar sind oder manuell extrahiert werden muessen."
ins.column_dimensions["A"].width = 130

wb.save(output_path)
print(output_path)
