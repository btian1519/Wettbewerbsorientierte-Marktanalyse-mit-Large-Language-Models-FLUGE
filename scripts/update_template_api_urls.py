from pathlib import Path

from openpyxl import load_workbook


PATH = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"


UPDATES = {
    "Preisniveau je Strecke/Zeitslot": {
        4: "https://developers.amadeus.com/self-service/category/flights/api-doc/flight-offers-search",
        5: "Ja",
        6: "hoch",
        7: "",
        8: "Posted fares != transaction fares",
        9: "Multi-horizon snapshots D-60/D-30/D-14/D-7",
        10: "https://duffel.com/docs/api",
        11: "Ja",
        12: "hoch",
        13: "",
        14: "Schema/coverage differs by provider",
        15: "Field harmonization + robust median",
        22: "2024-01 bis laufend (API polling)",
        29: "Price_proxy = robust_median(posted_fares over providers x booking_horizons)",
        30: "Posted fare differs from paid fare; provider schema heterogeneity",
        31: "Cross-provider median + horizon panel + outlier trimming",
        32: "MVP",
        33: "API-first sources, no manual UI scraping",
    },
    "Such- und Aufmerksamkeitsindikatoren": {
        4: "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/avia_par_be",
        5: "Ja",
        6: "hoch",
        7: "",
        8: "Monthly aggregation hides intra-month shocks",
        9: "Combine with high-frequency flight-state feeds",
        10: "https://opensky-network.org/apidoc/rest.html",
        11: "Teilweise",
        12: "hoch",
        13: "Demand_proxy from active flights and movements",
        14: "Coverage bias by ADS-B receiver density",
        15: "Cross-check against official monthly totals",
        22: "2023-01 bis laufend",
        29: "Demand_proxy = w1*normalized_passengers + w2*active_flight_index",
        30: "Different sampling frequency and spatial coverage",
        31: "Temporal alignment + calibration to official totals",
        32: "MVP",
        33: "Use ETL join by route-month and date",
    },
    "Puenktlichkeit und Annullierungen": {
        4: "https://aviationstack.com/documentation",
        5: "Ja",
        6: "hoch",
        7: "",
        8: "Free tier limits and endpoint throttling",
        9: "Retry/backoff + cached snapshots",
        10: "https://airlabs.co/docs",
        11: "Ja",
        12: "hoch",
        13: "",
        14: "Provider-specific status taxonomy",
        15: "Unified status mapping table",
        22: "2024-01 bis laufend (provider dependent)",
        29: "Ops_proxy = f(on_time_rate, cancel_rate, delay_quantiles)",
        30: "Status labels differ across providers",
        31: "Common event schema + provider fixed effects",
        32: "MVP",
        33: "Primary source class switched to API docs/endpoints",
    },
}


def main() -> None:
    wb = load_workbook(PATH)
    ws = wb["Datenquellen"]

    for row in range(2, ws.max_row + 1):
        info = ws.cell(row, 2).value
        if info in UPDATES:
            for col, value in UPDATES[info].items():
                ws.cell(row, col).value = value

    try:
        wb.save(PATH)
        print("Updated API-first rows in", PATH)
    except PermissionError:
        out = str(Path(PATH).with_name("Datenquellen_Template_FlightScope_api.xlsx"))
        wb.save(out)
        print("Primary file locked, wrote updated copy to", out)


if __name__ == "__main__":
    main()
