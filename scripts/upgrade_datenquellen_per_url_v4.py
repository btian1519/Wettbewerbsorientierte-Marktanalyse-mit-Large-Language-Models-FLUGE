from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

src = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"
out = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope_v4.xlsx"

wb = load_workbook(src)
old_ws = wb["Datenquellen"]

old_rows = []
for r in range(2, old_ws.max_row + 1):
    old_rows.append([old_ws.cell(r, c).value for c in range(1, old_ws.max_column + 1)])

idx = wb.sheetnames.index("Datenquellen")
del wb["Datenquellen"]
ws = wb.create_sheet("Datenquellen", idx)

headers = [
    "Modul",
    "Informationspunkt",
    "Pflichtfelder (mindestens)",
    "URL 1",
    "Direkt messbar URL 1? (Ja/Nein/Teilweise)",
    "Relevanz URL 1 (hoch/mittel/niedrig)",
    "URL 2",
    "Direkt messbar URL 2? (Ja/Nein/Teilweise)",
    "Relevanz URL 2 (hoch/mittel/niedrig)",
    "URL 3",
    "Direkt messbar URL 3? (Ja/Nein/Teilweise)",
    "Relevanz URL 3 (hoch/mittel/niedrig)",
    "URL 4",
    "Direkt messbar URL 4? (Ja/Nein/Teilweise)",
    "Relevanz URL 4 (hoch/mittel/niedrig)",
    "Abgedeckter Zeitraum",
    "Proxy-Definition (falls nicht direkt messbar)",
    "Bias-/Risiko-Hinweis",
    "Empfohlene Korrektur",
    "Prioritaet (MVP/P2/P3)",
    "Bemerkungen",
]

for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="00549F", end_color="00549F", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for i, row in enumerate(old_rows, start=2):
    modul, info, pflicht = row[0], row[1], row[2]
    u1, u2, u3, u4 = row[3], row[4], row[5], row[6]
    zeitraum, direkt_global = row[7], row[8]
    proxy, bias, korr, prio, bemerk = row[9], row[10], row[11], row[12], row[13]

    new_row = [
        modul, info, pflicht,
        u1, direkt_global if u1 else "", "",
        u2, direkt_global if u2 else "", "",
        u3, direkt_global if u3 else "", "",
        u4, direkt_global if u4 else "", "",
        zeitraum, proxy, bias, korr, prio, bemerk,
    ]

    for c, val in enumerate(new_row, 1):
        cell = ws.cell(i, c, val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

thin = Side(style="thin", color="D9D9D9")
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

widths = {1:22,2:34,3:46,4:28,5:18,6:20,7:28,8:18,9:20,10:28,11:18,12:20,13:28,14:18,15:20,16:18,17:34,18:26,19:28,20:16,21:24}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w

ws.row_dimensions[1].height = 44
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 56

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

wb.save(out)
print(out)
print("Rows:", ws.max_row, "Cols:", ws.max_column)
