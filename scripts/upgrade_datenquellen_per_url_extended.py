from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

path = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"

wb = load_workbook(path)
old_ws = wb["Datenquellen"]
old_rows = []
for r in range(2, old_ws.max_row + 1):
    old_rows.append([old_ws.cell(r, c).value for c in range(1, old_ws.max_column + 1)])

idx = wb.sheetnames.index("Datenquellen")
del wb["Datenquellen"]
ws = wb.create_sheet("Datenquellen", idx)

headers = [
    "Modul",
    "Informationspunkt",
    "Pflichtfelder (mindestens)",

    "URL 1",
    "Direkt messbar URL 1? (Ja/Nein/Teilweise)",
    "Relevanz URL 1 (hoch/mittel/niedrig)",
    "Proxy URL 1",
    "Bias/Risiko URL 1",
    "Korrektur URL 1",

    "URL 2",
    "Direkt messbar URL 2? (Ja/Nein/Teilweise)",
    "Relevanz URL 2 (hoch/mittel/niedrig)",
    "Proxy URL 2",
    "Bias/Risiko URL 2",
    "Korrektur URL 2",

    "URL 3",
    "Direkt messbar URL 3? (Ja/Nein/Teilweise)",
    "Relevanz URL 3 (hoch/mittel/niedrig)",
    "Proxy URL 3",
    "Bias/Risiko URL 3",
    "Korrektur URL 3",

    "URL 4",
    "Direkt messbar URL 4? (Ja/Nein/Teilweise)",
    "Relevanz URL 4 (hoch/mittel/niedrig)",
    "Proxy URL 4",
    "Bias/Risiko URL 4",
    "Korrektur URL 4",

    "Abgedeckter Zeitraum",
    "Proxy-Definition (Zeilenebene)",
    "Bias-/Risiko-Hinweis (Zeilenebene)",
    "Empfohlene Korrektur (Zeilenebene)",
    "Prioritaet (MVP/P2/P3)",
    "Bemerkungen",
]

for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="00549F", end_color="00549F", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# old schema (21 cols):
# 1 Modul,2 Info,3 Pflicht,
# 4 URL1,5 DM1,6 Rel1,
# 7 URL2,8 DM2,9 Rel2,
# 10 URL3,11 DM3,12 Rel3,
# 13 URL4,14 DM4,15 Rel4,
# 16 Zeitraum,17 ProxyGlobal,18 BiasGlobal,19 KorrGlobal,20 Prio,21 Bem
for r_idx, row in enumerate(old_rows, start=2):
    modul, info, pflicht = row[0], row[1], row[2]
    u1, dm1, rel1 = row[3], row[4], row[5]
    u2, dm2, rel2 = row[6], row[7], row[8]
    u3, dm3, rel3 = row[9], row[10], row[11]
    u4, dm4, rel4 = row[12], row[13], row[14]
    zeitraum, proxy_g, bias_g, korr_g, prio, bem = row[15], row[16], row[17], row[18], row[19], row[20]

    # Migrate: keep global fields; prefill URL1-specific fields from old global so no information is lost.
    new_row = [
        modul, info, pflicht,

        u1, dm1, rel1, proxy_g if u1 else "", bias_g if u1 else "", korr_g if u1 else "",
        u2, dm2, rel2, "", "", "",
        u3, dm3, rel3, "", "", "",
        u4, dm4, rel4, "", "", "",

        zeitraum, proxy_g, bias_g, korr_g, prio, bem,
    ]

    for c, val in enumerate(new_row, 1):
        cell = ws.cell(r_idx, c, val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

thin = Side(style="thin", color="D9D9D9")
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

widths = {
    1: 20, 2: 30, 3: 40,
    4: 24, 5: 16, 6: 18, 7: 22, 8: 22, 9: 22,
    10: 24, 11: 16, 12: 18, 13: 22, 14: 22, 15: 22,
    16: 24, 17: 16, 18: 18, 19: 22, 20: 22, 21: 22,
    22: 24, 23: 16, 24: 18, 25: 22, 26: 22, 27: 22,
    28: 18, 29: 30, 30: 26, 31: 28, 32: 16, 33: 22,
}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w

ws.row_dimensions[1].height = 46
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 58

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

wb.save(path)
print(path)
print("Rows:", ws.max_row, "Cols:", ws.max_column)
