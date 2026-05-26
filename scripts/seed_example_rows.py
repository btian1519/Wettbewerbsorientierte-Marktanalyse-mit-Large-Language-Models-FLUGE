from openpyxl import load_workbook

path = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"
wb = load_workbook(path)
ws = wb["Datenquellen"]

# Helper: find row by Informationspunkt (column B)
def find_row(info_name: str):
    for r in range(2, ws.max_row + 1):
        if (ws.cell(r, 2).value or "").strip() == info_name:
            return r
    return None

# Column mapping (v5-like 33 cols)
# 1 Modul,2 Informationspunkt,3 Pflichtfelder,
# URL1 group: 4..9
# URL2 group: 10..15
# URL3 group: 16..21
# URL4 group: 22..27
# 28 Zeitraum,29 Proxy-Zeilenebene,30 Bias-Zeilenebene,31 Korrektur-Zeilenebene,32 Prio,33 Bemerkungen

# Example set 1: Preisdaten
r = find_row("Preisniveau je Strecke/Zeitslot")
if r:
    ws.cell(r, 4).value = "https://www.google.com/travel/flights"
    ws.cell(r, 5).value = "Teilweise"
    ws.cell(r, 6).value = "hoch"
    ws.cell(r, 7).value = "Median fare per route/date from repeated snapshots"
    ws.cell(r, 8).value = "Posted fare != transaction fare"
    ws.cell(r, 9).value = "Collect multi-horizon snapshots (D-60/D-30/D-14/D-7)"

    ws.cell(r, 10).value = "https://www.skyscanner.com/"
    ws.cell(r, 11).value = "Teilweise"
    ws.cell(r, 12).value = "hoch"
    ws.cell(r, 13).value = "Route-date median from metasearch listings"
    ws.cell(r, 14).value = "Ranking/personalization effects"
    ws.cell(r, 15).value = "Use route-level median across times of day"

    ws.cell(r, 16).value = "https://www.kayak.com/flights"
    ws.cell(r, 17).value = "Teilweise"
    ws.cell(r, 18).value = "mittel"
    ws.cell(r, 19).value = "Fallback fare source for cross-check"
    ws.cell(r, 20).value = "Coverage differs by market"
    ws.cell(r, 21).value = "Treat as validation source, not primary"

    ws.cell(r, 28).value = "Rolling snapshot collection, weekly + event windows"
    ws.cell(r, 29).value = "Price_proxy = median(posted_fare across sources and horizons)"
    ws.cell(r, 30).value = "Observed fares may diverge from actual paid fares"
    ws.cell(r, 31).value = "Use robust medians, winsorization, and horizon controls"
    ws.cell(r, 32).value = "MVP"
    ws.cell(r, 33).value = "EXAMPLE filled by core team; route pilot FRA-LHR"

# Example set 2: Nachfrage-Proxies
r = find_row("Such- und Aufmerksamkeitsindikatoren")
if r:
    ws.cell(r, 4).value = "https://trends.google.com/trends/"
    ws.cell(r, 5).value = "Nein"
    ws.cell(r, 6).value = "hoch"
    ws.cell(r, 7).value = "Search index as demand-intent proxy"
    ws.cell(r, 8).value = "Interest != bookings"
    ws.cell(r, 9).value = "Use together with price and capacity features"

    ws.cell(r, 10).value = "https://www.google.com/travel/flights"
    ws.cell(r, 11).value = "Nein"
    ws.cell(r, 12).value = "mittel"
    ws.cell(r, 13).value = "Popularity signals from route exploration"
    ws.cell(r, 14).value = "Opaque platform logic"
    ws.cell(r, 15).value = "Use as auxiliary signal only"

    ws.cell(r, 28).value = "Weekly trend index"
    ws.cell(r, 29).value = "Demand_proxy = normalized search_index"
    ws.cell(r, 30).value = "Proxy may overreact to media/events"
    ws.cell(r, 31).value = "Smoothing + lag features + event controls"
    ws.cell(r, 32).value = "MVP"
    ws.cell(r, 33).value = "EXAMPLE filled by core team"

# Example set 3: Operative Qualitaet
r = find_row("Puenktlichkeit und Annullierungen")
if r:
    ws.cell(r, 4).value = "https://www.eurocontrol.int/Economics/DailyTrafficVariation-States.html"
    ws.cell(r, 5).value = "Teilweise"
    ws.cell(r, 6).value = "hoch"
    ws.cell(r, 7).value = "Use daily punctuality/operations indicators as ops proxy"
    ws.cell(r, 8).value = "Often state/airport-level, not pure airline-route"
    ws.cell(r, 9).value = "Map to route via airport-time matching"

    ws.cell(r, 10).value = "https://www.eurocontrol.int/performance"
    ws.cell(r, 11).value = "Teilweise"
    ws.cell(r, 12).value = "mittel"
    ws.cell(r, 13).value = "Performance dashboard proxy"
    ws.cell(r, 14).value = "Definition differences"
    ws.cell(r, 15).value = "Harmonize KPI definition (e.g., A14)"

    ws.cell(r, 28).value = "Daily/weekly operational metrics"
    ws.cell(r, 29).value = "Ops_proxy = f(OTP, cancellation, delay distribution)"
    ws.cell(r, 30).value = "Different reporting definitions across sources"
    ws.cell(r, 31).value = "Standardize KPI definitions before modeling"
    ws.cell(r, 32).value = "MVP"
    ws.cell(r, 33).value = "EXAMPLE filled by core team"

wb.save(path)
print(path)
