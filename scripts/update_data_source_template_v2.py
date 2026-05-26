from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

path = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"

wb = load_workbook(path)

# Replace / rebuild main sheet with enhanced schema
if "Datenquellen" in wb.sheetnames:
    del wb["Datenquellen"]
ws = wb.create_sheet("Datenquellen", 0)

headers = [
    "Modul",
    "Informationspunkt",
    "Pflichtfelder (mindestens)",
    "URL 1",
    "URL 2",
    "URL 3",
    "URL 4",
    "Abgedeckter Zeitraum",
    "Direkt messbar? (Ja/Nein)",
    "Proxy-Definition (falls Nein)",
    "Bias-/Risiko-Hinweis",
    "Empfohlene Korrektur",
    "Prioritaet (MVP/P2/P3)",
    "Bemerkungen",
]

rows = [
    [
        "Airline-Stammdaten",
        "Airline-Liste und Codes",
        "Airline-Name, IATA/ICAO, Land/Hub-Flughafen",
        "", "", "", "", "", "Ja", "", "Mapping-Fehler bei Codes", "Code-Validierung gegen IATA/ICAO", "MVP", "",
    ],
    [
        "Strecken- und Frequenzangebot",
        "Strecken und Frequenzen aller Airlines",
        "Abflugflughafen, Zielflughafen, Fluege pro Woche, Abflugzeitfenster, Flugzeugtyp",
        "", "", "", "", "", "Ja", "", "Saisonale Verschiebung", "Rolling Snapshot (woechentlich)", "MVP", "",
    ],
    [
        "Sitzplatz-/Kapazitaetsangebot",
        "Kapazitaetsdaten aller Airlines",
        "Strecke, Flugzeugtyp, Sitzplatzanzahl (oder ableitbare Felder)",
        "", "", "", "", "", "Teilweise", "Seat map je Flugzeugtyp als Standardwert", "Flottenkonfiguration variiert je Airline", "Airline-spezifische Seat-Maps priorisieren", "MVP", "",
    ],
    [
        "Preisdaten",
        "Preisniveau je Strecke/Zeitslot",
        "Crawling-Datum, Abflugdatum, Airline, Strecke, Mindestpreis/Medianpreis, Tarif-/Klassenhinweis",
        "", "", "", "", "", "Ja", "", "Posted fare != transaction fare", "Mehrere Buchungsvorlaeufe (D-60/D-30/D-14/D-7)", "MVP", "",
    ],
    [
        "Nachfrage-Proxies",
        "Such- und Aufmerksamkeitsindikatoren",
        "Keywords, Region, Zeitgranularitaet, Trendindex",
        "", "", "", "", "", "Nein", "Google Trends / Plattforminteresse als Nachfrage-Proxy", "Interest != actual bookings", "Mit Preis/Frequenz/Ops-Features gemeinsam verwenden", "MVP", "",
    ],
    [
        "Operative Qualitaet",
        "Puenktlichkeit und Annullierungen",
        "Airline/Strecke, OTP, Cancellation Rate, Delay-Verteilung, Definitionsbasis",
        "", "", "", "", "", "Teilweise", "Airport/State-level OTP als Ersatz", "Unterschiedliche KPI-Definitionen", "Definition harmonisieren (z. B. A14)", "MVP", "",
    ],
    [
        "Textdaten (Reviews)",
        "Service-Signale und Themen",
        "Review-Text, Datum, Plattform, Bewertung, Sprache",
        "", "", "", "", "", "Ja", "", "Selection bias (extreme Meinungen)", "Reweighting + Plattform-Fixed-Effects", "MVP", "",
    ],
    [
        "Kalender- und Eventdaten",
        "Exogene Nachfragefaktoren",
        "Feiertage, Messen, Sportevents, Datum, Stadt",
        "", "", "", "", "", "Ja", "", "Event-Effekt schwer isolierbar", "Dummy-Variablen + Lag-Features", "P2", "",
    ],
    [
        "Neue-Strecken-Potenzial",
        "O&D-Nachfragepotenzial",
        "Staedtepaar, Bevoelkerungs-/Tourismus-/Business-Indikatoren, Saisonalitaet",
        "", "", "", "", "", "Nein", "Regionale Makroindikatoren als O&D-Proxy", "Aggregationsbias", "Mehrere Quellen triangulieren", "P2", "",
    ],
    [
        "Wettbewerbsstruktur",
        "Wettbewerbslage pro Strecke",
        "Anzahl Carrier pro Strecke, Gesamtfrequenz/Gesamtsitze, Preisband",
        "", "", "", "", "", "Ja", "", "Carrier-Matching-Fehler", "Einheitliche Airline-ID erzwingen", "MVP", "",
    ],
    [
        "Flughafen-Restriktionen",
        "Operative und Slot-bezogene Grenzen",
        "Slot-Lage, Nachtflugbeschraenkungen, Terminal/Ground-Handling-Limits, regulatorische Auflagen",
        "", "", "", "", "", "Teilweise", "Proxy ueber publizierte Restriktionen", "Aktualitaetsrisiko", "Gueltigkeitsdatum zwingend speichern", "P2", "",
    ],
    [
        "Target Variable (empfohlen)",
        "Estimated Pax",
        "Seats_total, predicted_LF, route, timeslot, airline, date",
        "", "", "", "", "", "Nein", "Estimated_Pax = Seats_total * predicted_LF", "Fehler in LF uebertraegt sich", "Konfidenzintervall + Backtesting berichten", "MVP", "",
    ],
]

# Write headers
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="00549F", end_color="00549F", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Write rows
for r_idx, row in enumerate(rows, 2):
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# Borders
thin = Side(style="thin", color="D9D9D9")
for r in range(1, len(rows) + 2):
    for c in range(1, len(headers) + 1):
        ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Column widths
widths = {
    1: 24, 2: 34, 3: 56,
    4: 26, 5: 26, 6: 26, 7: 26,
    8: 20, 9: 17, 10: 34,
    11: 28, 12: 30, 13: 18, 14: 26,
}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w

ws.row_dimensions[1].height = 34
for r in range(2, len(rows) + 2):
    ws.row_dimensions[r].height = 58

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

# Add bias-check sheet
if "Bias_Check" in wb.sheetnames:
    del wb["Bias_Check"]
bc = wb.create_sheet("Bias_Check")

bc_headers = [
    "Bias-Typ",
    "Woran erkennbar?",
    "Risiko fuer Modell",
    "Korrektur in Datenpipeline",
    "Korrektur im Modell",
    "Monitoring-KPI",
]
bc_rows = [
    [
        "Selection Bias (Reviews)",
        "Uebergewicht extremer 1/5-Sterne Reviews",
        "Verzerrte Service-Signale",
        "Sampling ueber mehrere Plattformen, Duplikatfilter",
        "Reweighting + Plattform Fixed Effects",
        "Rating-Verteilung vs. Benchmark",
    ],
    [
        "Platform Bias",
        "Systematische Unterschiede je Plattform",
        "Nicht vergleichbare Sentiment-Scores",
        "Plattform als eigenes Feld speichern",
        "Platform Dummies / Hierarchical Model",
        "Score-Drift pro Plattform",
    ],
    [
        "Price Measurement Bias",
        "Unterschied zw. posted und transaction fare",
        "Fehlerhafte Preiselastizitaet",
        "Mehrere Advance Purchase Snapshots",
        "Robuste Elastizitaet + Szenarien",
        "MAPE der Preisfeatures",
    ],
    [
        "Temporal Bias",
        "Datenspikes in Ferien/Event-Zeiten",
        "Ueberschaetzung Peak-Nachfrage",
        "Saison/Event Flags",
        "Zeit-Fixed-Effects",
        "Fehler in Peak vs Off-Peak",
    ],
]

for c, h in enumerate(bc_headers, 1):
    cell = bc.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="0098A1", end_color="0098A1", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r_idx, row in enumerate(bc_rows, 2):
    for c_idx, val in enumerate(row, 1):
        cell = bc.cell(r_idx, c_idx, val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, w in {1: 30, 2: 32, 3: 28, 4: 34, 5: 30, 6: 24}.items():
    bc.column_dimensions[get_column_letter(c)].width = w
bc.row_dimensions[1].height = 30
for r in range(2, len(bc_rows) + 2):
    bc.row_dimensions[r].height = 52
bc.freeze_panes = "A2"

wb.save(path)
print(path)
