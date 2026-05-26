"""
FlightScope AI – PPT Generator
Takes the existing 'Branchenauswahl' file as base and appends new slides.
Style: RWTH Aachen corporate design (white BG, RWTH Blue headers).
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# ── RWTH Color Palette ─────────────────────────────────────────────────────
RWTH_BLUE   = RGBColor(0x00, 0x54, 0x9F)
RWTH_TEAL   = RGBColor(0x00, 0x61, 0x65)
RWTH_TEAL2  = RGBColor(0x00, 0x98, 0xA1)
RWTH_GREEN  = RGBColor(0x57, 0xAB, 0x27)
RWTH_ORANGE = RGBColor(0xF6, 0xA8, 0x00)
RWTH_RED    = RGBColor(0xCC, 0x07, 0x1E)
RWTH_LIGHT  = RGBColor(0x8E, 0xBA, 0xE5)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
GRAY_DARK   = RGBColor(0x40, 0x40, 0x40)
GRAY_MED    = RGBColor(0x70, 0x70, 0x70)
GRAY_LIGHT  = RGBColor(0xF0, 0xF2, 0xF5)
GRAY_BORDER = RGBColor(0xCC, 0xCC, 0xCC)

SRC = r'c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\20260429 Branchenauswahl.pptx'
DST = r'c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\FlightScope_AI_RWTH.pptx'

prs = Presentation(SRC)
W = prs.slide_width.inches   # 13.33
H = prs.slide_height.inches  # 7.5
BLANK = prs.slide_layouts[7]  # Inhalt_Aufzählung – cleanest blank-ish layout

# ── Low-level helpers ──────────────────────────────────────────────────────

def inches(v): return Inches(v)

def add_rect(slide, l, t, w, h, fill, line_color=None):
    shape = slide.shapes.add_shape(1, inches(l), inches(t), inches(w), inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape

def add_tb(slide, text, l, t, w, h,
           size=14, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
           italic=False, wrap=True):
    tb = slide.shapes.add_textbox(inches(l), inches(t), inches(w), inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb

def add_lines(slide, lines, l, t, w, h,
              size=13, color=GRAY_DARK, spacing=4, bullet=False):
    tb = slide.shapes.add_textbox(inches(l), inches(t), inches(w), inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(spacing)
        run = p.add_run()
        run.text = ("• " + line) if bullet else line
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return tb

# ── Slide components ────────────────────────────────────────────────────────

def header_bar(slide, title, subtitle=None, color=RWTH_BLUE):
    """Top bar with title."""
    add_rect(slide, 0, 0, W, 1.1, color)
    add_tb(slide, title, 0.4, 0.13, W - 0.8, 0.6,
           size=26, bold=True, color=WHITE)
    if subtitle:
        add_tb(slide, subtitle, 0.4, 0.72, W - 0.8, 0.38,
               size=12, color=RGBColor(0xCC, 0xDD, 0xFF), italic=True)

def footer_line(slide, n_total=9):
    add_rect(slide, 0, H - 0.28, W, 0.003, RWTH_BLUE)

def card(slide, l, t, w, h, accent=RWTH_BLUE, title=None, lines=None,
         title_size=13, line_size=12):
    """Filled card with optional accent bar, title, bullets."""
    add_rect(slide, l, t, w, h, GRAY_LIGHT, GRAY_BORDER)
    add_rect(slide, l, t, 0.07, h, accent)
    if title:
        add_tb(slide, title, l + 0.16, t + 0.1, w - 0.25, 0.42,
               size=title_size, bold=True, color=accent)
    if lines:
        add_lines(slide, lines, l + 0.16, t + (0.56 if title else 0.12),
                  w - 0.25, h - (0.7 if title else 0.2),
                  size=line_size, color=GRAY_DARK, spacing=3)

def flow_box(slide, label, desc, l, t, w, h, color=RWTH_BLUE):
    add_rect(slide, l, t, w, h, color)
    add_tb(slide, label, l + 0.1, t + 0.1, w - 0.2, 0.45,
           size=14, bold=True, color=WHITE)
    add_tb(slide, desc, l + 0.1, t + 0.58, w - 0.2, h - 0.7,
           size=11.5, color=WHITE, wrap=True)

def arrow(slide, l, t):
    add_tb(slide, "→", l, t, 0.45, 0.5,
           size=22, bold=True, color=GRAY_MED, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 2 — Literaturgrundlage
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "Literaturgrundlage", "Drei Quellen bilden die methodische Basis unseres Tools")

papers = [
    (RWTH_BLUE, "Porter – Competitive Strategy",
     "Methodischer Rahmen",
     ["Five Forces: strukturierte Analyse der Wettbewerbskraefte",
      "Luftfahrt: Hohe Rivalitaet, preistransparenter Markt",
      "Eintrittsbarrieren (Slots, Flotte, Regulierung)",
      "→ Zeigt: Warum LLM-Wettbewerbsanalyse hier wertvoll ist"]),
    (RWTH_TEAL, "IPA + Kano – Importance-Performance",
     "Dimensionsauswahl",
     ["Trennung: Basisfaktoren / Leistungsfaktoren / Begeisterung",
      "Luftfahrt-Basis: Sicherheit, Puenktlichkeit, klare Regeln",
      "Differenzierung: App-Qualitaet, kulante Umbuchung",
      "→ Zeigt: Welche Dimensionen wir priorisieren sollen"]),
    (RWTH_GREEN, "LLM-Cure – LLM-Based Competitor Analysis",
     "Technische Methode",
     ["LLM extrahiert Staerken/Schwaechen aus Reviews automatisch",
      "Standardisiertes Scoring ueber mehrere Wettbewerber",
      "Strukturierte JSON-Ausgabe pro Dimension + Belegzitate",
      "→ Zeigt: Wie wir die Analyse automatisieren"]),
]

for i, (color, title, sub, bullets) in enumerate(papers):
    x = 0.4 + i * 4.28
    add_rect(sl, x, 1.25, 4.08, 5.85, GRAY_LIGHT, GRAY_BORDER)
    add_rect(sl, x, 1.25, 4.08, 0.08, color)
    add_tb(sl, title, x + 0.15, 1.37, 3.8, 0.5, size=13.5, bold=True, color=color)
    add_tb(sl, sub, x + 0.15, 1.88, 3.8, 0.35, size=11, color=GRAY_MED, italic=True)
    add_lines(sl, bullets, x + 0.15, 2.28, 3.8, 4.5,
              size=12, color=GRAY_DARK, spacing=5, bullet=True)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 3 — Wettbewerbsdimensionen
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "Wettbewerbsdimensionen",
           "6 LLM-Analyse-Dimensionen + 2 operative Entscheidungsdimensionen", RWTH_TEAL)

dims_llm = [
    ("1", "Puenktlichkeit & Zuverlaessigkeit",    "OTP, Cancellation Rate, IrrOps-Handling"),
    ("2", "Preis & Gebuehrentransparenz",          "Preisklarheit, Zusatzkosten, Gepaeckregeln"),
    ("3", "Kundenservice & Stoerungsmanagement",   "Reaktionszeit, Problemloesung, Erreichbarkeit"),
    ("4", "Digitale Experience (App/Web)",         "App-Qualitaet, Self-Service, Online Check-in"),
    ("5", "Bord- & Reiseerlebnis",                 "Boarding, Sitzkomfort, Gepaeckabwicklung"),
    ("6", "Erstattungs- & Umbuchungsqualitaet",    "Rueckzahlungsgeschwindigkeit, Flexibilitaet"),
]
dims_ops = [
    ("7", "Preisniveau je Strecke/Zeitslot",      "Ticketpreise, Preiselastizitaet, Wettbewerberpreise"),
    ("8", "Nachfrage- & Passagierindikatoren",     "Suchtrends, Buchungsproxies, Sitzangebote, Auslastung"),
]

for i, (num, name, desc) in enumerate(dims_llm):
    col = i % 2
    row = i // 2
    x = 0.4 + col * 6.25
    y = 1.3 + row * 1.38
    add_rect(sl, x, y, 6.0, 1.2, GRAY_LIGHT, GRAY_BORDER)
    add_rect(sl, x, y, 0.5, 1.2, RWTH_TEAL)
    add_tb(sl, num, x, y + 0.28, 0.5, 0.55,
           size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_tb(sl, name, x + 0.62, y + 0.08, 5.25, 0.45,
           size=13, bold=True, color=RWTH_TEAL)
    add_tb(sl, desc, x + 0.62, y + 0.56, 5.25, 0.5,
           size=11, color=GRAY_MED)

# Ops dimensions
add_rect(sl, 0.4, 5.42, 12.53, 0.05, RWTH_ORANGE)
add_tb(sl, "Erweiterung fuer operative Kapazitaetsentscheidungen:",
       0.4, 5.52, 8.0, 0.38, size=12, bold=True, color=RWTH_ORANGE)
for i, (num, name, desc) in enumerate(dims_ops):
    x = 0.4 + i * 6.25
    y = 5.96
    add_rect(sl, x, y, 6.0, 1.1, RGBColor(0xFF, 0xF4, 0xDC), RWTH_ORANGE)
    add_rect(sl, x, y, 0.5, 1.1, RWTH_ORANGE)
    add_tb(sl, num, x, y + 0.22, 0.5, 0.55,
           size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_tb(sl, name, x + 0.62, y + 0.06, 5.25, 0.42,
           size=13, bold=True, color=RWTH_ORANGE)
    add_tb(sl, desc, x + 0.62, y + 0.52, 5.25, 0.45,
           size=11, color=GRAY_DARK)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 4 — Tool-Konzept: Workflow
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "Tool-Konzept: FlightScope AI",
           "End-to-End-Workflow: Von Rohdaten zur Kapazitaetsentscheidung", RWTH_BLUE)

steps = [
    (RWTH_BLUE,   "① Input",
     "Airlines, Strecke\nZeitraum, Quellen"),
    (RWTH_TEAL,   "② Datensammlung",
     "Reviews + Preiszeitreihen\n+ Nachfrage-Proxies"),
    (RWTH_TEAL2,  "③ LLM-Analyse",
     "Aspekt-Extraktion\nScoring + Belegzitate"),
    (RWTH_GREEN,  "④ Forecast / OR",
     "Nachfrageprognose\nSupply-Gap-Berechnung"),
    (RWTH_ORANGE, "⑤ Entscheidung",
     "Add / Hold / Reduce\nMachtyp + Zeitslot"),
]

bw = 2.18
bh = 2.2
for i, (color, label, desc) in enumerate(steps):
    x = 0.35 + i * (bw + 0.28)
    flow_box(sl, label, desc, x, 1.3, bw, bh, color)
    if i < 4:
        arrow(sl, x + bw + 0.02, 1.3 + bh / 2 - 0.25)

# Description row
descs = [
    "Zielairlines + Strecken definieren, Zeitraum & Quellen auswaehlen",
    "Reviews (Trustpilot/App Store), Ticketpreise, Suchtrend-Daten",
    "LLM strukturiert Freitext: Scores 1-10, Sentiment, Zitate je Dim.",
    "Demand Pressure + Supply Gap + Profitability Proxy berechnen",
    "Regelbasierte Empfehlung + Confidence Score + Top-3-Begruendung",
]
for i, desc in enumerate(descs):
    x = 0.35 + i * (bw + 0.28)
    add_tb(sl, desc, x, 3.65, bw, 1.1,
           size=11, color=GRAY_DARK, wrap=True)

# BYOK note
add_rect(sl, 0.35, 4.95, W - 0.7, 1.1, GRAY_LIGHT, GRAY_BORDER)
add_rect(sl, 0.35, 4.95, 0.07, 1.1, RWTH_BLUE)
add_tb(sl, "Software-Architektur: BYOK + OpenRouter Fallback",
       0.55, 5.05, 6.0, 0.42, size=13, bold=True, color=RWTH_BLUE)
add_lines(sl, [
    "BYOK (Bring Your Own Key): Nutzer gibt eigenen API-Key ein (Google / OpenRouter-kompatibel)",
    "Fallback auf OpenRouter: Automatischer Wechsel wenn kein gueltiger Key verfuegbar → stabile Demo",
    "Engine-Unabhaengigkeit: selbe Pipeline, austauschbares Modell-Backend",
], 0.55, 5.5, W - 0.9, 0.9, size=12, color=GRAY_DARK, spacing=2)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 5 — UI-Design
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "UI-Design: Dashboard-Aufbau",
           "Sidebar + Analysebereich + 5 Ergebnis-Tabs", RWTH_TEAL)

# Sidebar
add_rect(sl, 0.35, 1.25, 3.1, 5.85, GRAY_LIGHT, GRAY_BORDER)
add_rect(sl, 0.35, 1.25, 3.1, 0.08, RWTH_TEAL)
add_tb(sl, "Sidebar", 0.5, 1.38, 2.8, 0.42, size=14, bold=True, color=RWTH_TEAL)
add_lines(sl, [
    "API & Engine Settings",
    "  Key-Eingabe (Google / OpenRouter)",
    "  Modell-Auswahl",
    "  Test Connection",
    "",
    "Airlines-Auswahl (4-8)",
    "Strecke (Origin – Destination)",
    "Zeitraum & Quellenfilter",
    "",
    "▶  Run Analysis",
], 0.5, 1.85, 2.8, 5.0, size=12, color=GRAY_DARK, spacing=2)

# Tabs
tab_data = [
    (RWTH_BLUE,   "Overview",           "KPI-Karten\n+ Radar-Chart"),
    (RWTH_TEAL,   "Dim. Ranking",       "Ranking je\nDimension"),
    (RWTH_TEAL2,  "Head-to-Head",       "2 Airlines im\nDetail"),
    (RWTH_GREEN,  "Evidence",           "Belegzitate\n& Problemmuster"),
    (RWTH_ORANGE, "Capacity Planner",   "Add/Hold/Reduce\nMachtyp + Zeitslot"),
]
tw = 1.83
for i, (color, tab, content) in enumerate(tab_data):
    x = 3.65 + i * (tw + 0.12)
    # tab label bar
    add_rect(sl, x, 1.25, tw, 0.52, color)
    add_tb(sl, tab, x + 0.1, 1.3, tw - 0.15, 0.42,
           size=11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # content area
    add_rect(sl, x, 1.77, tw, 5.33, RGBColor(0xFA, 0xFA, 0xFC), GRAY_BORDER)
    add_tb(sl, content, x + 0.1, 2.2, tw - 0.2, 1.0,
           size=12, color=color, align=PP_ALIGN.CENTER, wrap=True)

# Capacity Planner detail in last tab
add_lines(sl, [
    "Entscheidungskarte:",
    "Add / Hold / Reduce",
    "",
    "Machtyp-Empfehlung",
    "Zeitslot-Empfehlung",
    "",
    "Top-3 Begruendungen",
    "+ Confidence Score",
], 3.65 + 4 * (tw + 0.12) + 0.12, 3.1,
   tw - 0.25, 3.8, size=11, color=GRAY_DARK, spacing=2)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 6 — MVP & Nutzen
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "MVP und Nutzen",
           "Minimum Viable Product: kleinste Version mit vollem Kernnutzen", RWTH_GREEN)

mvp_cols = [
    (RWTH_TEAL, "Wettbewerbsanalyse", [
        "Radarvergleich: 4-8 Airlines",
        "Evidenzbasierte Rankings je Dimension",
        "Head-to-Head Tiefenvergleich (2 Airlines)",
        "Sentiment-Zeittrend je Thema",
    ]),
    (RWTH_ORANGE, "Operative Entscheidung", [
        "Preis- & Nachfragemonitor je Strecke",
        "Regelbasierte Empfehlung:",
        "  Add / Hold / Reduce Flight",
        "Machtyp-Vorschlag (S/M/L Narrowbody)",
        "Empfohlenes Abflugzeitfenster",
    ]),
    (RWTH_BLUE, "Kernfrage, die wir beantworten", [
        "Soll Airline X auf Strecke r",
        "zum Zeitpunkt t die Frequenz",
        "erhoehen?",
        "",
        "Welche Flugzeuggroesse?",
        "(z. B. A320neo vs A321neo)",
        "",
        "Welches Zeitfenster ist optimal?",
    ]),
]
for i, (color, title, items) in enumerate(mvp_cols):
    x = 0.4 + i * 4.28
    card(sl, x, 1.25, 4.08, 5.85, color, title, items,
         title_size=14, line_size=12.5)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 7 — Capacity Planner Logik
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "Capacity Planner – Entscheidungslogik",
           "Input → Regel-Engine → Add / Hold / Reduce + Machtyp + Zeitslot", RWTH_ORANGE)

# Input box
add_rect(sl, 0.35, 1.25, 3.5, 5.85, GRAY_LIGHT, GRAY_BORDER)
add_rect(sl, 0.35, 1.25, 0.07, 5.85, RWTH_ORANGE)
add_tb(sl, "Input Features", 0.55, 1.35, 3.2, 0.45, size=13, bold=True, color=RWTH_ORANGE)
add_lines(sl, [
    "Preisniveau & -veraenderung",
    "Nachfrage-Proxies",
    "  (Suchtrend, Buchungs-Proxy)",
    "Wettbewerber-Kapazitaet",
    "Aktuelle Frequenz & Sitze",
    "OTP / Cancellation Rate",
    "Saisonalitaet & Kalender",
], 0.55, 1.88, 3.2, 4.8, size=12, color=GRAY_DARK, spacing=4)

arrow(sl, 3.85, 3.6)

# Rule Engine box
add_rect(sl, 4.3, 1.25, 4.7, 5.85, GRAY_LIGHT, GRAY_BORDER)
add_rect(sl, 4.3, 1.25, 0.07, 5.85, RWTH_BLUE)
add_tb(sl, "Regel-Engine (MVP, erklaerbar)", 4.5, 1.35, 4.4, 0.45,
       size=13, bold=True, color=RWTH_BLUE)
steps_eng = [
    "Schritt 1: Demand Pressure Score",
    "Schritt 2: Supply Gap Score",
    "Schritt 3: Profitability Proxy",
    "Schritt 4: Risikoabschlag (OTP/Service)",
    "Schritt 5: Final Score → Entscheidung",
    "",
    "Logik:",
    "  Demand hoch + Gap positiv + Risiko ok",
    "    → Add Flight",
    "  Mittlere Werte → Hold",
    "  Demand schwach / Risiko hoch → Reduce",
]
add_lines(sl, steps_eng, 4.5, 1.88, 4.4, 4.8, size=12, color=GRAY_DARK, spacing=3)

arrow(sl, 9.0, 3.6)

# Output box
add_rect(sl, 9.45, 1.25, 3.55, 5.85, GRAY_LIGHT, GRAY_BORDER)
add_rect(sl, 9.45, 1.25, 0.07, 5.85, RWTH_GREEN)
add_tb(sl, "Ausgabe", 9.65, 1.35, 3.25, 0.45, size=13, bold=True, color=RWTH_GREEN)
add_lines(sl, [
    "Entscheidung:",
    "  Add / Hold / Reduce",
    "",
    "Machtyp:",
    "  Mittl. Nachfrage → S-Narrowbody",
    "  Hohe Nachfrage/Peak → M/L-NB",
    "  Unsicher → erst Freq. testen",
    "",
    "Zeitslot:",
    "  Peak (Nachfrage + OTP stabil)",
    "  Off-Peak (preiselastisch)",
    "",
    "Top-3 Begruendungen",
    "+ Confidence Score",
], 9.65, 1.88, 3.25, 4.8, size=12, color=GRAY_DARK, spacing=3)

footer_line(sl)

# ══════════════════════════════════════════════════════════════════════════
# NEW SLIDE 8 — Machbarkeit & Next Steps
# ══════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header_bar(sl, "Machbarkeit und naechste Schritte",
           "Ja – in drei Phasen umsetzbar. Phase 1 im Kurs lieferbar.", RWTH_BLUE)

phases = [
    (RWTH_TEAL, "Phase 1 – Kurs-Demo", [
        "Reviews + Preise + Nachfrage-Proxies integrieren",
        "Regelbasiertes Add/Hold/Reduce implementieren",
        "Machtyp & Zeitslot-Heuristiken einbauen",
        "Streamlit-Dashboard mit allen 5 Tabs",
        "Demo: 2 Airlines x 2 Strecken",
    ], "Dieses Semester"),
    (RWTH_BLUE, "Phase 2 – Mittelfristig", [
        "Zeitreihenprognose je Strecke/Zeitslot",
        "Backtesting gegen historische Daten",
        "Automatisierte Datenpipeline",
        "Mehrsprachige Review-Quellen",
    ], "Nach dem Kurs"),
    (RWTH_ORANGE, "Phase 3 – Produktionsnah", [
        "Netzwerkweite Optimierung",
        "Flotten- & Slot-Constraints",
        "Szenarioanalyse (Fuel, Saison)",
        "Wettbewerber-Reaktionsmodell",
    ], "Langfristig"),
]
for i, (color, title, items, tag) in enumerate(phases):
    x = 0.4 + i * 4.28
    add_rect(sl, x, 1.25, 4.08, 5.1, GRAY_LIGHT, GRAY_BORDER)
    add_rect(sl, x, 1.25, 4.08, 0.08, color)
    add_tb(sl, title, x + 0.15, 1.38, 3.6, 0.45, size=14, bold=True, color=color)
    add_rect(sl, x + 0.15, 1.88, 1.5, 0.35, color)
    add_tb(sl, tag, x + 0.18, 1.92, 1.44, 0.28, size=10, bold=True, color=WHITE)
    add_lines(sl, items, x + 0.15, 2.33, 3.78, 3.7,
              size=12.5, color=GRAY_DARK, spacing=4, bullet=True)

# Action items
add_rect(sl, 0.4, 6.52, W - 0.8, 0.07, RWTH_BLUE)
add_tb(sl, "Naechste Schritte (sofort):", 0.4, 6.65, 3.5, 0.4,
       size=12, bold=True, color=RWTH_BLUE)
add_lines(sl, [
    "Scope festlegen (6-7 Airlines + 2 Pilotstrecken)",
    "Dimensionsschema & Prompt-Template finalisieren",
    "Erste Datenpipeline aufbauen (Reviews + Preise)",
    "Klickbare Demo bis zur naechsten Session liefern",
], 4.0, 6.65, 9.0, 0.75, size=11.5, color=GRAY_DARK, spacing=0)

footer_line(sl)

# ── Save ───────────────────────────────────────────────────────────────────
prs.save(DST)
print(f"Saved: {DST}")
print(f"Total slides: {len(prs.slides)}")
