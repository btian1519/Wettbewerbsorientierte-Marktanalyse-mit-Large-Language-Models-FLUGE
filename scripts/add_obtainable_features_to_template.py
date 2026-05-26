from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

path = r"c:\AAA_RWTH\Wettbewerbsorientierte Marktanalyse mit Large Language Models-FLUGE\Datenquellen_Template_FlightScope.xlsx"
wb = load_workbook(path)

ws = wb["Datenquellen"]

thin = Side(style="thin", color="D9D9D9")

new_rows = [
    [
        "Zeitfenster-Qualitaet",
        "Abflugzeit-Attraktivitaet",
        "Peak/Off-Peak Flag, red-eye Flag, business-friendly time window",
        "", "", "", "", "", "Ja", "", "Zeitzonen-/Sommerzeit-Effekte", "Zeit in lokale Slot-Kategorien mappen", "MVP", "",
    ],
    [
        "Zeitfenster-Qualitaet",
        "Schedule-Stabilitaet",
        "Anzahl timetable changes pro route/time window",
        "", "", "", "", "", "Teilweise", "Aenderungsrate aus wiederholten Snapshots", "Unvollstaendige Historie", "Woechentliche Snapshot-Versionierung", "P2", "",
    ],
    [
        "Preisstruktur",
        "Gesamtreisekosten-Proxy",
        "Fare + baggage fee + seat fee + change/refund penalty (wenn verfuegbar)",
        "", "", "", "", "", "Teilweise", "Total trip cost proxy bei fehlenden Komponenten", "Nicht alle Zusatzgebuehren oeffentlich", "Fehlende Komponenten als NA + Sensitivitaetsanalyse", "P2", "",
    ],
    [
        "Operative Robustheit",
        "Recovery-Faehigkeit bei Disruption",
        "Anzahl alternativer Fluege am selben Tag, mittlere Reaccommodation-Zeit",
        "", "", "", "", "", "Teilweise", "Alternative-flight count als Recovery-Proxy", "Nicht alle Umbuchungen beobachtbar", "Proxy klar kennzeichnen", "P2", "",
    ],
    [
        "Airport Experience",
        "Bodenprozess-Proxy",
        "security wait proxy, terminal load proxy, baggage wait proxy (falls verfuegbar)",
        "", "", "", "", "", "Nein", "Airport congestion proxy via time-of-day + airport ops data", "Direkte Messung selten", "Nur als schwaches Feature nutzen", "P3", "",
    ],
    [
        "Textdaten (Reviews)",
        "Review-Credibility und Themenintensitaet",
        "credibility label, topic intensity (baggage, delay, refund, service)",
        "", "", "", "", "", "Ja", "", "Selection bias + Bot/Spam Risiko", "De-duplication + credibility scoring + reweighting", "MVP", "",
    ],
    [
        "Wettbewerbsdynamik",
        "Relative Preis- und Angebotsluecke",
        "price_gap_vs_competitors, seat_share, frequency_gap",
        "", "", "", "", "", "Ja", "", "Falsches Carrier-Matching", "Einheitliche carrier IDs und route keys", "MVP", "",
    ],
]

start_row = ws.max_row + 1
for i, row in enumerate(new_rows):
    r = start_row + i
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Adjust row heights for new rows
for r in range(start_row, ws.max_row + 1):
    ws.row_dimensions[r].height = 58

# Keep filter current
ws.auto_filter.ref = f"A1:N{ws.max_row}"

# Add / refresh Feature_Catalog sheet
if "Feature_Catalog" in wb.sheetnames:
    del wb["Feature_Catalog"]
fc = wb.create_sheet("Feature_Catalog")

fc_headers = [
    "Feature_Name",
    "Typ",
    "Ebene",
    "Definition",
    "Quelle (Beispiel)",
    "Direkt/Proxy",
    "Bias-Risiko",
    "Empfohlene Behandlung",
    "Prioritaet",
]

catalog_rows = [
    ["weekly_frequency", "numeric", "route-week", "Anzahl Fluege pro Woche je Airline/Route/Timeslot", "Airline/airport schedules", "Direkt", "niedrig", "standardisieren", "MVP"],
    ["seats_total", "numeric", "route-week", "Gesamtsitze = Summe(Fluege*Seats je Typ)", "schedule + aircraft seat map", "Proxy", "mittel", "airline-spezifische seat map", "MVP"],
    ["price_median_d14", "numeric", "route-date", "Medianpreis 14 Tage vor Abflug", "OTA fare snapshots", "Direkt", "mittel", "mehrere advance windows", "MVP"],
    ["price_gap_vs_comp", "numeric", "route-date", "Relativer Preisabstand zum Wettbewerber-Median", "computed", "Direkt", "mittel", "winsorize", "MVP"],
    ["otp_rate", "numeric", "airline-route-month", "On-time performance rate", "regulator/airport stats", "Teilweise", "mittel", "A14 definition harmonisieren", "MVP"],
    ["cancel_rate", "numeric", "airline-route-month", "Anteil annullierter Fluege", "regulator/airport stats", "Teilweise", "mittel", "missingness report", "MVP"],
    ["search_trend_index", "numeric", "route-week", "Interesseindex fuer route-keywords", "Google Trends", "Proxy", "hoch", "nur mit Core Features zusammen nutzen", "MVP"],
    ["event_intensity", "numeric", "city-week", "Event/holiday Nachfrageimpuls", "public event calendars", "Proxy", "mittel", "event dummies + lag", "P2"],
    ["redeye_flag", "binary", "flight", "1 falls red-eye Flug", "schedule times", "Direkt", "niedrig", "local time conversion", "MVP"],
    ["schedule_change_rate", "numeric", "route-week", "Anteil geaenderter Flugzeiten", "snapshot diffs", "Proxy", "mittel", "versioned snapshots", "P2"],
    ["recovery_alt_flights", "numeric", "route-day", "Anzahl alternativer Fluege am selben Tag", "schedule graph", "Proxy", "mittel", "route-day graph build", "P2"],
    ["review_topic_delay", "numeric", "airline-week", "LLM topic intensity: delay", "reviews + LLM", "Proxy", "hoch", "platform FE + reweighting", "MVP"],
    ["review_topic_refund", "numeric", "airline-week", "LLM topic intensity: refund", "reviews + LLM", "Proxy", "hoch", "credibility score weighting", "MVP"],
    ["review_credibility_score", "numeric", "review", "Text- und accountbasierte Glaubwuerdigkeit", "reviews metadata", "Proxy", "hoch", "low-score downweight", "MVP"],
]

# Header style
for c, h in enumerate(fc_headers, 1):
    cell = fc.cell(1, c, h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="57AB27", end_color="57AB27", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for r_idx, row in enumerate(catalog_rows, 2):
    for c_idx, val in enumerate(row, 1):
        cell = fc.cell(r_idx, c_idx, val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

widths = {1: 26, 2: 10, 3: 16, 4: 44, 5: 28, 6: 12, 7: 14, 8: 32, 9: 10}
for c, w in widths.items():
    fc.column_dimensions[get_column_letter(c)].width = w

fc.row_dimensions[1].height = 28
for r in range(2, fc.max_row + 1):
    fc.row_dimensions[r].height = 44

fc.freeze_panes = "A2"
fc.auto_filter.ref = f"A1:I{fc.max_row}"

wb.save(path)
print(path)
print("Datenquellen rows:", ws.max_row)
print("Feature_Catalog rows:", fc.max_row)
